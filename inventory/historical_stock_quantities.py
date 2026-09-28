"""Read original purchase quantities for display, never for available stock."""

from collections import defaultdict
from decimal import Decimal

from openpyxl import load_workbook

from .legacy_stock_import import (
    MISSING_IDS,
    SOURCE_SHEET,
    decimal,
    normalized_description,
    text,
)


IDENTITY_FIELDS = ("Description", "Category", "Comp Type", "UOM")
REQUIRED_FIELDS = {"Comp ID", "Qty", "In Inv", *IDENTITY_FIELDS}


def stock_identity(values):
    """Keep different parts with the same old ID in separate groups."""
    code = text(values.get("Comp ID")).upper()
    if code in MISSING_IDS:
        return None
    fields = tuple(normalized_description(values.get(key)) for key in IDENTITY_FIELDS)
    if not fields[0]:
        return None
    return (code, *fields)


def source_quantities(filename):
    """Return (totals by component identity, identity by Excel row, invalid keys)."""
    workbook = load_workbook(filename, read_only=True, data_only=True)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing sheet: {SOURCE_SHEET}")
        rows = workbook[SOURCE_SHEET].iter_rows(values_only=True)
        next(rows, None)  # The title row is not a column header.
        headers = [text(cell) for cell in next(rows)]
        if REQUIRED_FIELDS - set(headers):
            raise ValueError("Use the original Stock Verification workbook with Qty and In Inv columns.")
        positions = {header: idx for idx, header in enumerate(headers) if header}
        totals = defaultdict(lambda: Decimal("0"))
        by_row = {}
        invalid = set()
        for excel_row, cells in enumerate(rows, start=3):
            values = {header: cells[idx] if idx < len(cells) else None
                      for header, idx in positions.items()}
            identity = stock_identity(values)
            if identity is None:
                continue
            by_row[excel_row] = identity
            qty = decimal(values.get("Qty"))
            if qty is None or qty < 0 or qty != qty.to_integral_value():
                invalid.add(identity)
            else:
                totals[identity] += qty
        return totals, by_row, invalid
    finally:
        workbook.close()
